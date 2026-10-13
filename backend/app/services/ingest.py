"""Materials -> text sections -> chunks -> embeddings (+ concept tag per chunk) -> Chroma."""
import json
import re
from pathlib import Path
import numpy as np
from . import rag


def _pdf(p):
    from pypdf import PdfReader
    return [(f"p.{i+1}", pg.extract_text() or "") for i, pg in enumerate(PdfReader(p).pages)]


def _pptx(p):
    from pptx import Presentation
    return [(f"slide {i+1}", "\n".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame))
            for i, sl in enumerate(Presentation(p).slides)]


def _docx(p):
    from docx import Document
    out, cur = [], ["intro", []]
    for para in Document(p).paragraphs:
        if para.style.name.startswith("Heading"):
            out.append((cur[0], "\n".join(cur[1])))
            cur = [para.text, []]
        else:
            cur[1].append(para.text)
    out.append((cur[0], "\n".join(cur[1])))
    return out


def _nb(p):
    cells = json.loads(Path(p).read_text(encoding="utf-8"))["cells"]
    return [(f"cell {i+1}", "".join(c.get("source", ""))) for i, c in enumerate(cells)]


def _txt(p):
    return [("text", Path(p).read_text(encoding="utf-8", errors="ignore"))]


EXT = {".pdf": _pdf, ".pptx": _pptx, ".docx": _docx, ".ipynb": _nb, ".md": _txt, ".txt": _txt}
KIND = {".pdf": "PDF", ".pptx": "PowerPoint", ".docx": "Word", ".ipynb": "Notebook", ".md": "Markdown", ".txt": "Text"}


def chunk(text: str, size: int = 900) -> list[str]:
    out, cur = [], ""
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if cur and len(cur) + len(para) > size:
            out.append(cur)
            cur = ""
        cur = (cur + "\n\n" + para).strip()
        while len(cur) > size * 1.5:
            out.append(cur[:size])
            cur = cur[size:]
    if cur:
        out.append(cur)
    return out


def ingest(path: Path, material_id: int, concepts: list) -> int:
    docs, metas = [], []
    for sec, text in EXT[path.suffix.lower()](str(path)):
        for ch in chunk(text):
            docs.append(ch)
            metas.append({"material": path.name.split("_", 1)[-1], "material_id": material_id, "section": sec, "concept": ""})
    if not docs:
        return 0
    emb = np.array(rag.ef()(docs), dtype=float)
    if concepts:  # tag each chunk with its nearest course concept (cosine)
        cv = np.array(rag.ef()([f"{c.name}. {c.description}" for c in concepts]), dtype=float)
        sims = (emb / np.linalg.norm(emb, axis=1, keepdims=True)) @ (cv / np.linalg.norm(cv, axis=1, keepdims=True)).T
        for i, m in enumerate(metas):
            j = int(sims[i].argmax())
            m["concept"] = concepts[j].key if sims[i, j] >= 0.25 else ""
    rag.col().add(ids=[f"m{material_id}-{i}" for i in range(len(docs))], documents=docs, embeddings=emb.tolist(), metadatas=metas)
    return len(docs)
